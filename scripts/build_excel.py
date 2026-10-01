"""Construit la base Excel cote-auto.xlsx à partir du JSON produit par src/export.mjs.

    node src/export.mjs | python3 scripts/build_excel.py cote-auto.xlsx

Le fichier est entièrement régénéré à chaque passage du robot à partir de l'historique complet :
il s'alimente donc tout seul, mois après mois. Les colonnes d'analyse (variations, position dans
la fourchette, décote, jours en ligne...) sont des formules Excel, recalculées à l'ouverture.
"""

import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
BRASS = "A8742F"
INK = "1F1C18"
DIM = "6E655A"
LINE = "E1D8CA"
F_BASE = Font(name=FONT, size=10, color=INK)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color=INK)
F_DIM = Font(name=FONT, size=9, italic=True, color=DIM)
F_LINK = Font(name=FONT, size=10, color="1F5FA8", underline="single")
F_KPI = Font(name=FONT, size=14, bold=True, color=BRASS)
FILL_HEAD = PatternFill("solid", fgColor=BRASS)
FILL_BUY = PatternFill("solid", fgColor="E3F1EA")
FILL_WATCH = PatternFill("solid", fgColor="F7EBD6")
BORDER = Border(bottom=Side(style="thin", color=LINE))

EUR = '#,##0 "€";-#,##0 "€";"-"'
KM = '#,##0 "km";;"-"'
PCT = '0.0%'
TREND = '[Color10]+0.0%;[Red]-0.0%;0.0%'
DELTA_PTS = '[Color10]+0.0;[Red]-0.0;0.0'
INT = '0'
DAYS = '0 "j"'
DATE = 'dd/mm/yyyy'


def d(s):
    """Date ISO -> date Excel (ou None)."""
    return date.fromisoformat(s[:10]) if s else None


def table(ws, row, headers, rows, formats=None, widths=None, formulas=None, links=None):
    """Écrit un tableau à partir de la ligne `row` : en-tête stylé, filtres, volets figés.

    headers  : [(titre, clé du dict ou None pour une formule)]
    formulas : {index de colonne: fonction(r) -> formule pour la ligne r}
    links    : {index de colonne: clé de l'URL} — la cellule affiche « Voir » et pointe sur l'URL
    """
    formats = formats or {}
    formulas = formulas or {}
    links = links or {}
    for c, (title, _) in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=title)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    for i, item in enumerate(rows):
        r = row + 1 + i
        for c, (_, key) in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=c)
            if c in formulas:
                cell.value = formulas[c](r)
            elif c in links:
                url = item.get(links[c])
                if url:
                    cell.value = "Voir"
                    cell.hyperlink = url
                    cell.font = F_LINK
                    cell.border = BORDER
                    continue
            elif key is not None:
                cell.value = item.get(key)
            cell.font = F_BASE
            cell.border = BORDER
            if c in formats:
                cell.number_format = formats[c]
    last = row + max(len(rows), 1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{last}"
    ws.freeze_panes = ws.cell(row=row + 1, column=2)
    for c, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return row + 1, row + len(rows)


def signal_colors(ws, col, first, last):
    if last < first:
        return
    rng = f"{col}{first}:{col}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("achat",{col}{first}))'], fill=FILL_BUY))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("surveiller",{col}{first}))'], fill=FILL_WATCH))


def build(data, out):
    wb = Workbook()

    # ------------------------------------------------------------------ Lisez-moi
    ws = wb.active
    ws.title = "Lisez-moi"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "COTE — base de suivi des cotes"
    ws["A1"].font = F_TITLE
    lines = [
        ("Généré le", d(data["genere_le"])),
        ("Dernier relevé", d(data["dernier_releve"]) or "aucun relevé automatique pour l'instant"),
        ("Dashboard", data["dashboard"]),
        ("", ""),
        ("Fonctionnement", "Ce fichier est régénéré automatiquement à chaque relevé (chaque lundi et le 1er du mois) à partir de tout l'historique."),
        ("", "Il est aussi joint au récap mensuel par mail. Toute modification faite à la main ici sera écrasée : enregistre une copie pour annoter."),
        ("", ""),
        ("Onglets", ""),
        ("Analyse du mois", "Situation actuelle de chaque véhicule : score, signal, prix, tendances, offre, ventes. Indicateurs clés en haut."),
        ("Historique mensuel", "Une ligne par véhicule et par mois (dernier relevé du mois). Variations vs mois précédent calculées par formule."),
        ("Médianes par mois", "Tableau croisé véhicule × mois des prix médians (formules sur l'onglet Historique mensuel)."),
        ("Indice base 100", "Évolution de la médiane depuis le premier mois suivi (100 = premier mois), avec graphique."),
        ("Relevés", "Tous les relevés bruts, y compris l'amorce manuelle de septembre 2026."),
        ("Annonces", "Chaque annonce avec tout ce que fournit LeBonCoin : version et finition exactes, prix, ancien prix, estimation de prix LeBonCoin (min/max, positionnement), km, année, mise en circulation, carburant, boîte, puissance, couleur, entretien, historique, vendeur pro/particulier, localisation, photos, description ; durée en ligne et baisses de prix."),
        ("Historique des prix", "Chaque prix relevé pour chaque annonce, avec l'évolution d'un relevé à l'autre."),
        ("Par année-modèle", "Prix moyen, plus bas, plus haut et km moyen par véhicule et par année-modèle (annonces en vente)."),
        ("Par kilométrage", "Mêmes indicateurs par tranche de kilométrage : mesure la décote liée au km."),
        ("Bonnes affaires", "Annonces du dernier relevé au moins 15 % sous la cote attendue pour leur kilométrage."),
        ("Véhicules", "Liste des véhicules suivis ou retirés et leurs critères de recherche."),
        ("", ""),
        ("Définitions", ""),
        ("Médiane", "Prix médian des annonces retenues (après filtres : pièces, épaves, autres versions, doublons, prix aberrants)."),
        ("Cote à km réf.", "Prix attendu à un kilométrage de référence fixe (régression prix/km) : compare les mois sans biais de kilométrage."),
        ("Position fourchette", "0 % = au plancher, 100 % = au plafond de prix observé (ou saisi pendant les 3 premiers mois)."),
        ("Score", "Sur 100 : rareté 20 %, désirabilité 20 %, « dernier de » 10 %, proximité du plancher 20 %, momentum 15 %, marché 15 %."),
        ("Ventes", "Annonces disparues entre deux relevés. Une annonce peut aussi être retirée sans vente : c'est une estimation."),
        ("Avertissement", "Outil de suivi de marché, pas un conseil d'achat."),
    ]
    for i, (a, b) in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=a).font = Font(name=FONT, size=10, bold=True, color=INK)
        cell = ws.cell(row=i, column=2, value=b)
        cell.font = F_BASE
        if isinstance(b, date):
            cell.number_format = DATE
            cell.alignment = Alignment(horizontal="left")
        if isinstance(b, str) and b.startswith("http"):
            cell.hyperlink = b
            cell.font = F_LINK

    # ------------------------------------------------------------------ Analyse du mois
    ws = wb.create_sheet("Analyse du mois")
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"Analyse au {d(data['genere_le']).strftime('%d/%m/%Y')}"
    ws["A1"].font = F_TITLE
    syn = data["synthese"]
    H = 8
    first, last = H + 1, H + max(len(syn), 1)
    kpis = [
        ("Véhicules suivis", f"=COUNTA(A{first}:A{last})", INT),
        ("Score moyen", f'=IFERROR(AVERAGEIF(N{first}:N{last},">0",C{first}:C{last}),0)', "0.0"),
        ("Annonces actives", f"=SUM(N{first}:N{last})", INT),
        ("Ventes sur 30 jours", f"=SUM(O{first}:O{last})", INT),
        ("Signaux d'achat", f'=COUNTIF(D{first}:D{last},"Signal d\'achat")', INT),
        ("Bonnes affaires", f"=SUM(Q{first}:Q{last})", INT),
    ]
    for i, (label, formula, fmt) in enumerate(kpis):
        col = 1 + i * 2 if i < 3 else 1 + (i - 3) * 2
        r = 3 if i < 3 else 5
        ws.cell(row=r, column=col, value=label).font = F_DIM
        cell = ws.cell(row=r + 1, column=col, value=formula)
        cell.font = F_KPI
        cell.number_format = fmt
        cell.alignment = Alignment(horizontal="left")
    headers = [
        ("Véhicule", "vehicule"), ("Catégorie", "categorie"), ("Score /100", "score"), ("Signal", "signal"),
        ("Médiane", "mediane"), ("Plancher", "plancher"), ("Plafond", "plafond"), ("Position fourchette", None),
        ("Tendance 1 mois", "tendance_1m"), ("Tendance 3 mois", "tendance_3m"), ("Tendance 12 mois", "tendance_12m"),
        ("Cote à km réf.", "cote_km_ref"), ("Km réf.", "km_ref"), ("Annonces", "annonces"), ("Vendues (30 j)", "vendues_30j"),
        ("Délai de vente", "delai_vente"), ("Bonnes affaires", "bonnes_affaires"), ("Rareté", "rarete"),
        ("Désirabilité", "desirabilite"), ("Dernier de", "dernier_de"), ("Proximité plancher", "proximite"),
        ("Momentum", "momentum"), ("Marché", "marche"), ("Dernier relevé", "dernier_releve"), ("LeBonCoin", None),
    ]
    for row in syn:
        row["dernier_releve"] = d(row["dernier_releve"])
    fmts = {3: "0.0", 5: EUR, 6: EUR, 7: EUR, 8: PCT, 9: TREND, 10: TREND, 11: TREND, 12: EUR, 13: KM,
            14: INT, 15: INT, 16: DAYS, 17: INT, 18: INT, 19: INT, 20: INT, 21: INT, 22: INT, 23: INT, 24: DATE}
    table(ws, H, headers, syn, formats=fmts,
          widths=[34, 30, 10, 15, 12, 12, 12, 12, 12, 12, 12, 13, 12, 10, 10, 10, 10, 9, 11, 10, 11, 11, 9, 13, 11],
          formulas={8: lambda r: f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r}),ISNUMBER(G{r}),G{r}>F{r}),(E{r}-F{r})/(G{r}-F{r}),"")'},
          links={25: "lien"})
    signal_colors(ws, "D", first, last)
    analyse_ws, analyse_first, analyse_last = ws, first, last

    if syn:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Score par véhicule"
        chart.style = 2
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 100
        chart.legend = None
        chart.add_data(Reference(ws, min_col=3, min_row=H, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        chart.height = max(7, 0.55 * len(syn) + 2)
        chart.width = 18
        chart.x_axis.scaling.orientation = "maxMin"
        ws.add_chart(chart, f"A{last + 3}")

    # ------------------------------------------------------------------ Historique mensuel
    ws = wb.create_sheet("Historique mensuel")
    men = data["mensuel"]
    headers = [
        ("Mois", "mois"), ("Véhicule", "vehicule"), ("Médiane", "mediane"), ("Évol. médiane vs mois préc.", None),
        ("P25", "p25"), ("P75", "p75"), ("Cote à km réf.", "cote_km_ref"), ("Évol. cote km réf.", None),
        ("Annonces", "annonces"), ("Nouvelles annonces", "nouvelles"), ("Vendues", "vendues"), ("Délai de vente", "delai_vente"),
        ("Relevés dans le mois", "releves"), ("Score /100", "score"), ("Évol. score (pts)", None), ("Signal", "signal"),
    ]
    same = lambda r: f"$B{r}=$B{r - 1}"
    evol = lambda col: lambda r: f'=IF(AND({same(r)},ISNUMBER({col}{r}),ISNUMBER({col}{r - 1})),{col}{r}/{col}{r - 1}-1,"")'
    f1, l1 = table(ws, 1, headers, men,
                   formats={3: EUR, 4: TREND, 5: EUR, 6: EUR, 7: EUR, 8: TREND, 9: INT, 10: INT, 11: INT, 12: DAYS, 13: INT, 14: "0.0", 15: DELTA_PTS},
                   widths=[10, 34, 12, 14, 12, 12, 13, 13, 10, 11, 9, 10, 10, 10, 11, 15],
                   formulas={4: evol("C"), 8: evol("G"),
                             15: lambda r: f'=IF(AND({same(r)},ISNUMBER(N{r}),ISNUMBER(N{r - 1})),N{r}-N{r - 1},"")'})
    signal_colors(ws, "P", f1, l1)
    hist_last = max(l1, 2)

    # ------------------------------------------------------------------ Médianes par mois
    months = sorted({m["mois"] for m in men})
    vehicles = []
    for m in men:
        if m["vehicule"] not in vehicles:
            vehicles.append(m["vehicule"])
    ws = wb.create_sheet("Médianes par mois")
    ws.cell(row=1, column=1, value="Véhicule")
    for j, month in enumerate(months, start=2):
        ws.cell(row=1, column=j, value=month)
    for c in range(1, len(months) + 2):
        cell = ws.cell(row=1, column=c)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
    rng = lambda col: f"'Historique mensuel'!${col}$2:${col}${hist_last}"
    for i, v in enumerate(vehicles, start=2):
        ws.cell(row=i, column=1, value=v).font = F_BASE
        for j in range(2, len(months) + 2):
            col = get_column_letter(j)
            cell = ws.cell(row=i, column=j,
                           value=f'=IFERROR(AVERAGEIFS({rng("C")},{rng("B")},$A{i},{rng("A")},{col}$1),"")')
            cell.font = F_BASE
            cell.number_format = EUR
    ws.column_dimensions["A"].width = 34
    for j in range(2, len(months) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 12
    ws.freeze_panes = "B2"

    # ------------------------------------------------------------------ Indice base 100
    ws = wb.create_sheet("Indice base 100")
    ws["A1"] = "Évolution de la médiane, base 100 au premier mois suivi de chaque véhicule"
    ws["A1"].font = Font(name=FONT, size=12, bold=True, color=INK)
    ws["A2"] = ("Valeurs calculées par le robot à partir de l'onglet Historique mensuel (cases vides = pas de relevé ce mois-là, "
                "pour que le graphique ne les compte pas comme des zéros).")
    ws["A2"].font = F_DIM
    R0 = 4
    ws.cell(row=R0, column=1, value="Mois")
    for j, v in enumerate(vehicles, start=2):
        ws.cell(row=R0, column=j, value=v)
    for c in range(1, len(vehicles) + 2):
        cell = ws.cell(row=R0, column=c)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[R0].height = 42
    base = {}
    by_key = {(m["vehicule"], m["mois"]): m["mediane"] for m in men if m["mediane"]}
    for m in men:
        if m["mediane"] and m["vehicule"] not in base:
            base[m["vehicule"]] = m["mediane"]
    for i, month in enumerate(months, start=R0 + 1):
        ws.cell(row=i, column=1, value=month).font = F_BASE
        for j, v in enumerate(vehicles, start=2):
            val = by_key.get((v, month))
            if val and base.get(v):
                cell = ws.cell(row=i, column=j, value=round(val / base[v] * 100, 1))
                cell.font = F_BASE
                cell.number_format = "0.0"
    ws.column_dimensions["A"].width = 12
    for j in range(2, len(vehicles) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.freeze_panes = ws.cell(row=R0 + 1, column=2)
    if months and vehicles:
        chart = LineChart()
        chart.title = "Indice de la médiane (100 = premier mois suivi)"
        chart.style = 2
        chart.y_axis.title = "Indice"
        chart.height = 11
        chart.width = 26
        chart.display_blanks = "gap"
        chart.add_data(Reference(ws, min_col=2, max_col=len(vehicles) + 1, min_row=R0, max_row=R0 + len(months)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=R0 + 1, max_row=R0 + len(months)))
        ws.add_chart(chart, f"A{R0 + len(months) + 3}")

    # ------------------------------------------------------------------ Relevés
    ws = wb.create_sheet("Relevés")
    rel = data["releves"]
    for r in rel:
        r["date"] = d(r["date"])
    headers = [
        ("Date", "date"), ("Véhicule", "vehicule"), ("Source", "source"), ("Annonces", "annonces"), ("Min", "min"),
        ("P25", "p25"), ("Médiane", "mediane"), ("P75", "p75"), ("Max", "max"), ("Km médian", "km_median"),
        ("Cote à km réf.", "cote_km_ref"), ("Km réf.", "km_ref"), ("Nouvelles", "nouvelles"), ("Vendues", "vendues"),
        ("Délai de vente", "delai_vente"), ("Part de pros", "part_pro"), ("Annonces écartées", "ecartees"),
    ]
    table(ws, 1, headers, rel,
          formats={1: DATE, 4: INT, 5: EUR, 6: EUR, 7: EUR, 8: EUR, 9: EUR, 10: KM, 11: EUR, 12: KM, 13: INT, 14: INT, 15: DAYS, 16: PCT, 17: INT},
          widths=[11, 34, 9, 10, 11, 11, 11, 11, 11, 12, 13, 12, 10, 9, 10, 10, 11])

    # ------------------------------------------------------------------ Annonces
    ws = wb.create_sheet("Annonces")
    ann = data["annonces"]
    for a in ann:
        for k in ("publiee_le", "premiere_vue", "derniere_vue", "partie_le"):
            a[k] = d(a[k])
    # (titre, clé, format, largeur) — None comme clé = colonne calculée par formule.
    cols = [
        ("Véhicule", "vehicule", None, 28), ("Titre", "titre", None, 38), ("Version", "version", None, 30),
        ("Finition", "finition", None, 12), ("Statut", "statut", None, 10), ("Publiée le", "publiee_le", DATE, 11),
        ("Première vue", "premiere_vue", DATE, 11), ("Dernière vue", "derniere_vue", DATE, 11), ("Partie le", "partie_le", DATE, 11),
        ("Jours en ligne", None, DAYS, 10), ("Prix initial", "prix_initial", EUR, 11), ("Prix actuel", "prix", EUR, 11),
        ("Baisse de prix", None, PCT, 10), ("Nb de baisses", "nb_baisses", INT, 9), ("Ancien prix affiché", "ancien_prix", EUR, 11),
        ("Estimation LBC min", "estimation_min", EUR, 11), ("Estimation LBC max", "estimation_max", EUR, 11),
        ("Écart vs estimation LBC", None, TREND, 11), ("Positionnement LBC", "positionnement", None, 16),
        ("Km", "km", KM, 12), ("Année", "annee", INT, 8), ("Mise en circulation", "mise_en_circulation", None, 11),
        ("Carburant", "carburant", None, 10), ("Boîte", "boite", None, 11), ("Puissance DIN (ch)", "puissance_din", INT, 10),
        ("Puissance fiscale (CV)", "puissance_fiscale", INT, 10), ("Couleur", "couleur", None, 11), ("Portes", "portes", INT, 7),
        ("Type", "type", None, 11), ("Entretien", "entretien", None, 30), ("Historique", "historique", None, 18),
        ("Vendeur", "vendeur", None, 11), ("Ville", "ville", None, 16), ("Code postal", "code_postal", None, 10),
        ("Département", "departement", None, 16), ("Région", "region", None, 16), ("Photos", "photos", INT, 8),
        ("Remontée payante", "remontee", None, 10), ("Urgent", "urgent", None, 8),
        ("Description (extrait)", "description", None, 60), ("Autres caractéristiques", "autres", None, 50),
        ("Annonce", None, None, 9), ("Id", "id", None, 13),
    ]
    C = {title: get_column_letter(i) for i, (title, *_rest) in enumerate(cols, start=1)}
    c = lambda title, r: f"{C[title]}{r}"
    formulas = {
        # Jours en ligne : depuis la publication si connue, sinon depuis la première observation.
        "Jours en ligne": lambda r: (f'=IF(OR(ISNUMBER({c("Publiée le", r)}),ISNUMBER({c("Première vue", r)})),'
                                     f'IF(ISNUMBER({c("Partie le", r)}),{c("Partie le", r)},{c("Dernière vue", r)})'
                                     f'-IF(ISNUMBER({c("Publiée le", r)}),{c("Publiée le", r)},{c("Première vue", r)}),"")'),
        "Baisse de prix": lambda r: (f'=IF(AND(ISNUMBER({c("Prix initial", r)}),ISNUMBER({c("Prix actuel", r)}),{c("Prix initial", r)}>0),'
                                     f'1-{c("Prix actuel", r)}/{c("Prix initial", r)},"")'),
        # Prix de l'annonce comparé au milieu de la fourchette estimée par LeBonCoin.
        "Écart vs estimation LBC": lambda r: (f'=IF(AND(ISNUMBER({c("Prix actuel", r)}),ISNUMBER({c("Estimation LBC min", r)}),'
                                              f'ISNUMBER({c("Estimation LBC max", r)})),{c("Prix actuel", r)}/'
                                              f'(({c("Estimation LBC min", r)}+{c("Estimation LBC max", r)})/2)-1,"")'),
    }
    idx = {title: i for i, (title, *_rest) in enumerate(cols, start=1)}
    table(ws, 1, [(t, k) for t, k, _, _ in cols], ann,
          formats={i: f for i, (_, _, f, _) in enumerate(cols, start=1) if f},
          widths=[w for *_, w in cols],
          formulas={idx[t]: f for t, f in formulas.items()},
          links={idx["Annonce"]: "lien"})
    ann_last = max(len(ann) + 1, 2)
    A = lambda title: f"Annonces!${C[title]}$2:${C[title]}${ann_last}"

    # ------------------------------------------------------------------ Historique des prix
    ws = wb.create_sheet("Historique des prix")
    px = data["prix"]
    for h in px:
        h["date"] = d(h["date"])
    headers = [("Véhicule", "vehicule"), ("Id annonce", "id"), ("Titre", "titre"), ("Date", "date"), ("Prix", "prix"),
               ("Évolution vs prix précédent", None), ("Annonce", None)]
    table(ws, 1, headers, px, formats={4: DATE, 5: EUR, 6: TREND}, widths=[28, 13, 40, 11, 11, 14, 9],
          formulas={6: lambda r: f'=IF(AND($B{r}=$B{r - 1},ISNUMBER(E{r - 1}),E{r - 1}>0),E{r}/E{r - 1}-1,"")'},
          links={7: "lien"})

    # ------------------------------------------------------------------ Par année-modèle / par kilométrage
    # Calculés par formules sur l'onglet Annonces (annonces en vente) : se mettent à jour seuls.
    def segment_sheet(title, rows, key_header, crit):
        ws = wb.create_sheet(title)
        ws["A1"] = f"Annonces en vente, par véhicule et {key_header.lower()} (formules sur l'onglet Annonces)"
        ws["A1"].font = F_DIM
        headers = [("Véhicule", "vehicule"), (key_header, "cle"), ("Annonces", None), ("Prix moyen", None),
                   ("Prix le plus bas", None), ("Prix le plus haut", None), ("Km moyen", None)]
        base = lambda r: f'{A("Véhicule")},$A{r},{A("Statut")},"En vente",{crit(r)}'
        table(ws, 2, headers, rows, formats={3: INT, 4: EUR, 5: EUR, 6: EUR, 7: KM}, widths=[34, 16, 10, 12, 13, 13, 12],
              formulas={3: lambda r: f"=COUNTIFS({base(r)})",
                        4: lambda r: f'=IFERROR(AVERAGEIFS({A("Prix actuel")},{base(r)}),"")',
                        5: lambda r: f'=IF(C{r}>0,_xlfn.MINIFS({A("Prix actuel")},{base(r)}),"")',
                        6: lambda r: f'=IF(C{r}>0,_xlfn.MAXIFS({A("Prix actuel")},{base(r)}),"")',
                        7: lambda r: f'=IFERROR(AVERAGEIFS({A("Km")},{base(r)}),"")'})
        return ws

    en_vente = [a for a in ann if a["statut"] == "En vente"]
    years = sorted({(a["vehicule"], a["annee"]) for a in en_vente if a["annee"]})
    segment_sheet("Par année-modèle", [{"vehicule": v, "cle": y} for v, y in years], "Année-modèle",
                  lambda r: f'{A("Année")},$B{r}')
    BANDS = [("< 100 000 km", 0, 100000), ("100 000 – 150 000 km", 100000, 150000),
             ("150 000 – 200 000 km", 150000, 200000), ("≥ 200 000 km", 200000, 10**7)]
    band_of = {label: (lo, hi) for label, lo, hi in BANDS}
    km_rows = []
    for v in sorted({a["vehicule"] for a in en_vente if a["km"]}):
        for label, lo, hi in BANDS:
            if any(a["vehicule"] == v and a["km"] and lo <= a["km"] < hi for a in en_vente):
                km_rows.append({"vehicule": v, "cle": label, "lo": lo, "hi": hi})
    ws_km = segment_sheet("Par kilométrage", km_rows, "Tranche de km",
                          lambda r: f'{A("Km")},">="&$H{r},{A("Km")},"<"&$I{r}')
    # Bornes de chaque tranche (colonnes H-I, utilisées par les formules).
    ws_km.cell(row=2, column=8, value="Km min").font = F_DIM
    ws_km.cell(row=2, column=9, value="Km max (exclu)").font = F_DIM
    for i, row in enumerate(km_rows, start=3):
        ws_km.cell(row=i, column=8, value=row["lo"]).font = F_DIM
        ws_km.cell(row=i, column=9, value=row["hi"]).font = F_DIM

    # ------------------------------------------------------------------ Bonnes affaires
    ws = wb.create_sheet("Bonnes affaires")
    aff = data["affaires"]
    for a in aff:
        a["releve"] = d(a["releve"])
    headers = [
        ("Véhicule", "vehicule"), ("Titre", "titre"), ("Prix", "prix"), ("Cote attendue", "cote_attendue"), ("Décote", None),
        ("Km", "km"), ("Année", "annee"), ("Ville", "ville"), ("Annonce", None), ("Relevé", "releve"),
    ]
    table(ws, 1, headers, aff,
          formats={3: EUR, 4: EUR, 5: PCT, 6: KM, 7: INT, 10: DATE},
          widths=[30, 44, 11, 13, 9, 12, 8, 16, 9, 11],
          formulas={5: lambda r: f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),D{r}>0),1-C{r}/D{r},"")'},
          links={9: "lien"})
    if not aff:
        ws["A3"] = "Aucune annonce nettement sous la cote au dernier relevé."
        ws["A3"].font = F_DIM

    # ------------------------------------------------------------------ Véhicules
    ws = wb.create_sheet("Véhicules")
    veh = data["vehicules"]
    for v in veh:
        v["ajoute_le"] = d(v["ajoute_le"])
    headers = [
        ("Véhicule", "vehicule"), ("Identifiant", "id"), ("Statut", "statut"), ("Catégorie", "categorie"),
        ("Recherche", "recherche"), ("Années", "annees"), ("Mots obligatoires", "obligatoires"), ("Variantes", "variantes"),
        ("Mots exclus", "exclus"), ("Rareté", "rarete"), ("Dernier de", "dernier_de"), ("Désirabilité", "desirabilite"),
        ("Ajouté le", "ajoute_le"), ("Notes", "notes"), ("LeBonCoin", None),
    ]
    table(ws, 1, headers, veh, formats={10: INT, 11: INT, 12: INT, 13: DATE},
          widths=[34, 22, 9, 30, 20, 11, 16, 20, 28, 8, 9, 10, 11, 60, 11], links={15: "lien"})

    # Colonnes de l'analyse calculées sur l'onglet Annonces (ajoutées une fois celui-ci écrit).
    ws = analyse_ws
    sel = lambda r: f'{A("Véhicule")},$A{r},{A("Statut")},"En vente"'
    extra = [
        ("Part de pros (en vente)", lambda r: f'=IFERROR(COUNTIFS({sel(r)},{A("Vendeur")},"Pro")/COUNTIFS({sel(r)},{A("Vendeur")},"<>"),"")', PCT),
        ("Part avec carnet d'entretien", lambda r: f'=IFERROR(COUNTIFS({sel(r)},{A("Entretien")},"*arnet*")/COUNTIFS({sel(r)},{A("Entretien")},"<>"),"")', PCT),
        ("Écart moyen vs estimation LBC", lambda r: f'=IFERROR(AVERAGEIFS({A("Écart vs estimation LBC")},{sel(r)}),"")', TREND),
        ("Jours en ligne moyens", lambda r: f'=IFERROR(AVERAGEIFS({A("Jours en ligne")},{sel(r)}),"")', DAYS),
    ]
    for j, (title, f, fmt) in enumerate(extra, start=26):
        cell = ws.cell(row=8, column=j, value=title)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = 12
        if syn:
            for r in range(analyse_first, analyse_last + 1):
                c = ws.cell(row=r, column=j, value=f(r))
                c.font = F_BASE
                c.number_format = fmt
                c.border = BORDER
    ws.auto_filter.ref = f"A8:{get_column_letter(25 + len(extra))}{analyse_last}"

    wb.calculation.fullCalcOnLoad = True
    wb.save(out)


if __name__ == "__main__":
    build(json.load(sys.stdin), sys.argv[1] if len(sys.argv) > 1 else "cote-auto.xlsx")
